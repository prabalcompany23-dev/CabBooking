"""Rebuild Phase-1.docx with 7 sections; tables copied verbatim from source."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(r"c:\Users\singhp8\OneDrive - Autodesk\Desktop\Business analyst career path")
TABLE_SOURCE = ROOT / "Phase-1-backup.docx"
BACKUP = ROOT / "Phase-1-backup.docx"
OUTPUT = ROOT / "Phase-1.docx"


def load_source_tables():
    doc = Document(TABLE_SOURCE)
    if len(doc.tables) != 3:
        raise RuntimeError(f"Expected 3 tables, found {len(doc.tables)}")
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)
    return tables


def add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(text, level=level)


def add_table_from_data(doc: Document, data: list[list[str]]):
    rows, cols = len(data), len(data[0])
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    for r_idx, row_data in enumerate(data):
        for c_idx, value in enumerate(row_data):
            table.rows[r_idx].cells[c_idx].text = value


def add_objective_section(doc: Document):
    add_heading(doc, "1. Objective", level=1)
    doc.add_paragraph(
        "Phase 1 defines the foundation for transitioning from approximately six years of "
        "Technical Sales experience (plus MBA Data Analytics) into a job-ready Business "
        "Analyst profile—not an entry-level fresher BA."
    )
    doc.add_paragraph(
        "Step 1 competency areas and their sequence are frozen. The exact course list is "
        "not frozen yet; Section 4 captures the current learning direction only."
    )
    doc.add_paragraph("Positioning:")
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(
        "Existing experience: Customer → Demo → Discussion → Pain Point → Feedback → "
        "Industry Challenge → Solution"
    )
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(
        "BA capability: Stakeholder → Problem Statement → Requirement → Analysis → "
        "Process → User Story → Acceptance Criteria → Prioritization → Agile Delivery → "
        "Validation → Business Outcome"
    )
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("MBA Data Analytics: Data → SQL → BI → Business Insight")
    doc.add_paragraph(
        "That combination is what we ultimately want to present to the market."
    )


def add_areas_section(doc: Document, priority_table: list[list[str]]):
    add_heading(doc, "2. Areas and Priority", level=1)
    doc.add_paragraph(
        "We finalized 11 competency areas, in this order. Step 1 = FROZEN."
    )
    add_table_from_data(doc, priority_table)


def add_matrix_section(doc: Document, matrix_table: list[list[str]]):
    add_heading(doc, "3. Key Topic-Skills Matrix (Proficiency)", level=1)
    doc.add_paragraph(
        "Proficiency targets, job-ready definitions, and syllabus keywords are frozen "
        "as documented below."
    )
    add_table_from_data(doc, matrix_table)


def add_learning_section(doc: Document):
    add_heading(doc, "4. Learning Resources (Recommended)", level=1)
    doc.add_paragraph("Learning Resources — Current Direction")
    doc.add_paragraph(
        "We have not yet frozen the exact course list. We have only established the approach."
    )
    doc.add_paragraph("Our preferred learning ecosystem will be:")

    blocks = [
        ("Business Analysis", [
            "Udemy / Coursera / LinkedIn Learning",
            "Potential primary options discussed:",
            "Mastering Business Analysis, Scrum and JIRA for Development — Udemy",
            "Business Analysis Fundamentals — Coursera",
            "Getting Started as a Business Analyst — LinkedIn Learning",
        ]),
        ("Agile/Scrum", [
            "Introduction to Agile Development and Scrum — Coursera",
            "Agile Business Analysis — LinkedIn Learning",
            "Udemy Agile/Scrum course",
            "Then PSM I",
        ]),
        ("Software/Product", [
            "Introduction to Software Product Management — University of Alberta, Coursera",
            "LinkedIn Learning product-management fundamentals",
        ]),
        ("SQL", [
            "SQL for Data Science — UC Davis, Coursera",
            "SQL Essential Training — LinkedIn Learning",
            "The Complete SQL Bootcamp: Go from Zero to Hero — Udemy",
        ]),
        ("Power BI", [
            "Data Analysis and Visualization with Power BI — Microsoft/Coursera",
            "Power BI Essential Training — LinkedIn Learning",
            "PL-300 later if desired",
        ]),
        ("Jira/Miro", [
            "Atlassian learning resources",
            "Miro learning resources",
            "YouTube/Udemy for practical usage",
        ]),
    ]
    for title, lines in blocks:
        add_heading(doc, title, level=2)
        for line in lines:
            doc.add_paragraph(line, style="List Bullet")
    doc.add_paragraph("Exact course selection is NOT frozen yet.")


def add_certifications_section(doc: Document, cert_table: list[list[str]]):
    add_heading(doc, "5. Certifications", level=1)
    add_table_from_data(doc, cert_table)


def add_placeholder_section(doc: Document, number: int, title: str):
    add_heading(doc, f"{number}. {title}", level=1)
    doc.add_paragraph("[To be completed.]")


def build():
    if not TABLE_SOURCE.exists():
        raise FileNotFoundError(TABLE_SOURCE)
    priority, matrix, certs = load_source_tables()
    matrix = deepcopy(matrix)
    certs = deepcopy(certs)
    priority = deepcopy(priority)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Phase 1 — Business Analyst Career Path", level=0)
    title.alignment = None

    add_objective_section(doc)
    add_areas_section(doc, priority)
    add_matrix_section(doc, matrix)
    add_learning_section(doc)
    add_certifications_section(doc, certs)
    add_placeholder_section(doc, 6, "Questionnaire")
    add_placeholder_section(doc, 7, "CV")

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Backup at {BACKUP}")


if __name__ == "__main__":
    build()
